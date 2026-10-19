import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

def build_sheet_resources(wb, font_family="Segoe UI"):
    ws = wb.create_sheet("Official Learning Resources")
    ws.views.sheetView[0].showGridLines = True
    
    font_title = Font(name=font_family, size=14, bold=True, color="1A365D")
    font_sub = Font(name=font_family, size=10, italic=True, color="4A5568")
    font_header = Font(name=font_family, size=9, bold=True, color="FFFFFF")
    font_data = Font(name=font_family, size=9, color="2D3748")
    font_bold = Font(name=font_family, size=9, bold=True, color="2D3748")
    font_link = Font(name=font_family, size=9, color="0969DA", underline="single")
    
    fill_header = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    border_thin = Border(
        left=Side(style='thin', color='CBD5E0'), right=Side(style='thin', color='CBD5E0'),
        top=Side(style='thin', color='CBD5E0'), bottom=Side(style='thin', color='CBD5E0')
    )

    # Title
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 18
    ws['A1'] = "CURATED OFFICIAL & AUTHORITATIVE AI/ML RESOURCES"
    ws['A1'].font = font_title
    ws['A2'] = "Verified Official Documentation, Cookbooks, Courses & Research Guides from Google, OpenAI, Meta, Hugging Face, PyTorch, and Leading Frameworks"
    ws['A2'].font = font_sub

    headers = [
        "Resource ID", "Organization", "Resource Title", "Official URL",
        "Learning Category", "Resource Type", "Level", "Prerequisites",
        "Recommended Phase", "Estimated Duration (hrs)", "Free / Paid Status",
        "Importance", "Learning Objectives", "Practical Exercise",
        "Completion Status", "Notes"
    ]
    
    header_row = 4
    ws.row_dimensions[header_row].height = 28
    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = border_thin
    ws.freeze_panes = "A5"

    resources = [
        # GOOGLE
        ("RES-001", "Google", "Google Machine Learning Crash Course", "https://developers.google.com/machine-learning/crash-course",
         "Machine Learning", "Course / Interactive", "Beginner", "Python basics", "Phase 2", 15, "Free", "Essential",
         "Learn foundational ML concepts, gradient descent, loss functions, classification, and neural network intuitions.",
         "Complete all embedded Colab exercises and model tuning tasks.", "Not Started", "Excellent visual explanations and interactive TensorFlow/Keras widgets."),
        
        ("RES-002", "Google", "Google Foundational ML & Data Prep Courses", "https://developers.google.com/machine-learning/data-prep",
         "Data Engineering", "Course / Guide", "Beginner to Intermediate", "Python, Statistics", "Phase 2", 10, "Free", "Recommended",
         "Master data transformation, feature scaling, numerical bucketing, categorical encoding, and handling missing values.",
         "Clean an uncurated customer dataset and transform features for regression modeling.", "Not Started", "Official Google guidance on real-world ML data hygiene."),
        
        ("RES-003", "Google", "Google Gemini API Documentation", "https://ai.google.dev/docs",
         "LLM APIs", "Official Documentation", "Intermediate", "REST APIs, Python", "Phase 5", 8, "Free / Freemium", "Essential",
         "Learn Gemini model families (Flash, Pro), multimodal inputs, system instructions, safety settings, and token limits.",
         "Build a Python client calling Gemini-1.5-Flash with custom safety thresholds and system prompt.", "Not Started", "Free tier offers generous rate limits for rapid prototyping."),
        
        ("RES-004", "Google", "Google Gemini API Cookbook", "https://github.com/google-gemini/cookbook",
         "LLM Applications", "Repository / Cookbook", "Intermediate", "Python, Gemini API", "Phase 5 & 6", 12, "Free", "Essential",
         "Practical recipe code for RAG, function calling, structured extraction, audio/video analysis, and document reasoning.",
         "Implement the official function calling and structured JSON output notebook recipes.", "Not Started", "Maintained directly by Google AI Developer Relations."),

        # OPENAI
        ("RES-005", "OpenAI", "OpenAI Developer Platform Documentation", "https://platform.openai.com/docs",
         "LLM APIs", "Official Documentation", "Beginner to Intermediate", "REST APIs, JSON", "Phase 5", 10, "Paid API", "Essential",
         "Understand Chat Completions, context windows, seed parameters, role system, streaming responses, and token usage.",
         "Create multi-turn chat wrapper handling token limits and streaming outputs to terminal.", "Not Started", "The industry benchmark API standard for commercial LLM integration."),
        
        ("RES-006", "OpenAI", "OpenAI Function Calling Guide", "https://platform.openai.com/docs/guides/function-calling",
         "AI Agents", "Official Guide", "Intermediate", "Python, JSON Schema", "Phase 7", 8, "Paid API", "Essential",
         "Learn strict tool schemas, parallel tool calls, tool choice parameters, and executing function responses.",
         "Build agent with 3 simulated business tools: CRM lookup, inventory query, discount calculator.", "Not Started", "Critical foundation for all agentic workflows."),
        
        ("RES-007", "OpenAI", "OpenAI Structured Outputs Guide", "https://platform.openai.com/docs/guides/structured-outputs",
         "LLM Applications", "Official Guide", "Intermediate", "Pydantic, JSON Schema", "Phase 5", 6, "Paid API", "Essential",
         "Master 100% reliable schema adherence with constrained decoding and native Pydantic integration.",
         "Extract validated lead BANT qualification objects directly from messy sales call transcripts.", "Not Started", "Eliminates need for retry loops due to malformed JSON."),
        
        ("RES-008", "OpenAI", "OpenAI Retrieval, Evals & Agent Building Guide", "https://platform.openai.com/docs/guides/evals",
         "Evaluation & Agents", "Guide / Framework", "Advanced", "Python, LLM APIs", "Phase 7 & 11", 12, "Free / Paid API", "Recommended",
         "Learn rigorous LLM evaluation methodologies, synthetic dataset creation, and agent benchmarking.",
         "Build an automated evaluation pipeline measuring sales agent response accuracy against 30 gold standards.", "Not Started", "Essential for verifying capability before selling to clients."),

        # META
        ("RES-009", "Meta", "Meta Llama Documentation & Model Card", "https://llama.meta.com/docs",
         "Open-Weight LLMs", "Official Documentation", "Intermediate to Advanced", "Linux, Python", "Phase 4 & 10", 10, "Free", "Essential",
         "Understand Llama-3 architecture, tokenizer nuances, prompt formatting, context window expansion, and licensing terms.",
         "Download and verify Llama-3-8B weights on Ubuntu Linux environment.", "Not Started", "Premier open-weight model family for local and cost-effective commercial deployment."),
        
        ("RES-010", "Meta", "Meta Llama Cookbook", "https://github.com/meta-llama/llama-cookbook",
         "LLM Engineering", "Repository / Cookbook", "Intermediate to Advanced", "Python, PyTorch", "Phase 6 & 10", 15, "Free", "Essential",
         "Production code for local inference with vLLM/Ollama, fine-tuning with PEFT, synthetic data, and RAG pipelines.",
         "Run end-to-end local RAG pipeline with Llama-3-8B and ChromaDB.", "Not Started", "Official Meta repository with battle-tested engineering notebooks."),
        
        ("RES-011", "Meta", "Meta Llama Fine-Tuning & Quantization Guide", "https://github.com/meta-llama/llama-recipes",
         "Fine-Tuning", "Cookbook / Guide", "Advanced", "PyTorch, CUDA, PEFT", "Phase 10", 14, "Free (Compute Req)", "Recommended",
         "Learn LoRA, QLoRA, FSDP distributed training, and quantizing Llama models for high-throughput inference.",
         "Fine-tune Llama-3-8B on an objection-handling dataset using QLoRA.", "Not Started", "Enables adapting open models to proprietary company sales playbooks."),

        # HUGGING FACE
        ("RES-012", "Hugging Face", "Hugging Face NLP Course", "https://huggingface.co/learn/nlp-course",
         "NLP & Transformers", "Interactive Course", "Beginner to Intermediate", "Python, PyTorch", "Phase 3 & 4", 25, "Free", "Essential",
         "Master Transformers from scratch, tokenizers, pipeline API, fine-tuning with Trainer, and evaluation metrics.",
         "Fine-tune a DistilBERT text classification model for customer email urgency detection.", "Not Started", "Widely recognized as the best practical deep dive into modern NLP."),
        
        ("RES-013", "Hugging Face", "Hugging Face Transformers Documentation", "https://huggingface.co/docs/transformers",
         "Transformers", "Official Documentation", "Intermediate to Advanced", "Python, PyTorch", "Phase 4", 20, "Free", "Essential",
         "Comprehensive API reference for model loading, quantization configurations (bitsandbytes, AWQ), generation parameters.",
         "Write custom generation loop with temperature, top_p, and stop token controls.", "Not Started", "The backbone library for open-source AI engineering."),

        # PYTORCH & FAST.AI
        ("RES-014", "PyTorch", "PyTorch Official Tutorials (Deep Learning with PyTorch)", "https://pytorch.org/tutorials",
         "Deep Learning", "Official Tutorials", "Beginner to Intermediate", "Python, Linear Algebra", "Phase 3", 20, "Free", "Essential",
         "Learn tensors, autograd, nn.Module, Dataset/DataLoader, optimizers, and model saving/loading.",
         "Build and train a 3-layer neural network from scratch on custom tabular sales data.", "Not Started", "The primary deep learning framework used in modern AI research and engineering."),
        
        ("RES-015", "fast.ai", "Practical Deep Learning for Coders (Part 1 & 2)", "https://course.fast.ai",
         "Deep Learning", "Video Course / Book", "Beginner to Intermediate", "Python", "Phase 3", 30, "Free", "Recommended",
         "Top-down, code-first introduction to deep learning, training dynamics, transfer learning, and model debugging.",
         "Train an image classifier and tabular prediction model using fastai/PyTorch.", "Not Started", "Exceptional intuitive pedagogy that demystifies neural networks."),

        # SCIKIT-LEARN & PYTHON FOUNDATIONS
        ("RES-016", "scikit-learn", "scikit-learn Official User Guide", "https://scikit-learn.org/stable/user_guide.html",
         "Machine Learning", "Official User Guide", "Beginner to Intermediate", "Python, NumPy", "Phase 2", 20, "Free", "Essential",
         "In-depth documentation of linear models, trees, ensembles, clustering, preprocessing pipelines, and cross-validation.",
         "Build a production Scikit-Learn Pipeline combining ColumnTransformer and RandomForestClassifier.", "Not Started", "The definitive reference for classical machine learning algorithms."),
        
        ("RES-017", "Python.org", "Python 3 Official Documentation & Typing Guide", "https://docs.python.org/3/",
         "Programming", "Official Documentation", "Beginner to Advanced", "None", "Phase 1", 15, "Free", "Essential",
         "Master data models, OOP, async io, typing module (TypeVar, Protocol), and exception hierarchies.",
         "Write an asynchronous task worker with typed protocol interfaces and custom exception handlers.", "Not Started", "Clean Python engineering is the single biggest differentiator of great AI engineers."),
        
        ("RES-018", "NumPy", "NumPy Official Documentation & Quickstart", "https://numpy.org/doc/stable/user/quickstart.html",
         "Scientific Computing", "Official Guide", "Beginner", "Python", "Phase 1 & 2", 8, "Free", "Essential",
         "Master multi-dimensional arrays, indexing, slicing, vectorization, linear algebra routines, and broadcasting.",
         "Implement cosine similarity and matrix multiplication from scratch using NumPy arrays.", "Not Started", "The foundation for all tensor computing and machine learning."),
        
        ("RES-019", "Pandas", "Pandas Official Documentation & User Guide", "https://pandas.pydata.org/docs/user_guide/",
         "Data Analysis", "Official Guide", "Beginner to Intermediate", "Python", "Phase 2", 12, "Free", "Essential",
         "DataFrames, Series, merging, indexing, time series, aggregation, and handling dirty business data.",
         "Load 50,000 CRM rows, clean timestamps, engineer lead velocity features, and export clean CSV.", "Not Started", "Essential tool for analyzing customer logs and training tabular models."),

        # AGENT & RAG FRAMEWORKS
        ("RES-020", "LangChain", "LangGraph Documentation & Agent Tutorials", "https://langchain-ai.github.io/langgraph/",
         "AI Agents", "Official Documentation", "Intermediate to Advanced", "Python, LLM APIs", "Phase 7 & 8", 18, "Free", "Essential",
         "Learn state graphs, cyclical agent loops, persistence, human-in-the-loop checkpoints, and multi-agent coordination.",
         "Build stateful sales qualification graph with conditional routing and human approval pauses.", "Not Started", "The industry standard framework for reliable, production-grade agent workflows."),
        
        ("RES-021", "CrewAI", "CrewAI Official Documentation", "https://docs.crewai.com/",
         "Multi-Agent", "Official Documentation", "Intermediate", "Python, LLM APIs", "Phase 8", 10, "Free", "Recommended",
         "Role-based autonomous multi-agent teams, task delegation, shared memory, and tools.",
         "Create a 3-agent team: Lead Researcher, Pitch Drafter, and QA Auditor.", "Not Started", "Intuitive role-playing abstraction for multi-agent collaboration."),

        # BACKEND, DATABASE & INFRASTRUCTURE
        ("RES-022", "FastAPI", "FastAPI Official Documentation", "https://fastapi.tiangolo.com/",
         "Backend Engineering", "Official Documentation", "Intermediate", "Python, Type hints", "Phase 11", 12, "Free", "Essential",
         "Async API design, Pydantic validation, dependency injection, OpenAPI documentation, and background tasks.",
         "Build production webhook API endpoint receiving messages, invoking agent graph, and returning async response.", "Not Started", "The standard high-performance Python web framework for AI APIs."),
        
        ("RES-023", "pgvector / PostgreSQL", "PostgreSQL & pgvector Official Documentation", "https://github.com/pgvector/pgvector",
         "Database & Vectors", "Official Documentation", "Intermediate", "SQL basics", "Phase 6 & 11", 8, "Free", "Essential",
         "Vector storage in relational DBs, IVFFlat and HNSW indexes, hybrid SQL + vector queries.",
         "Write SQL query combining metadata filtering (customer_industry='SaaS') with vector cosine similarity.", "Not Started", "Keeps your architecture simple by using one database for both relational data and vector embeddings."),
        
        ("RES-024", "Docker", "Docker Documentation & Container Best Practices", "https://docs.docker.com/",
         "DevOps & Deployment", "Official Documentation", "Intermediate", "Linux CLI", "Phase 11", 10, "Free", "Essential",
         "Dockerfiles, multi-stage builds, non-root user security, docker-compose orchestration, environment variables.",
         "Create multi-container docker-compose stack running FastAPI, PostgreSQL pgvector, and Redis.", "Not Started", "Ensures reproducibility and painless deployment to Render, Hetzner, or AWS."),
        
        ("RES-025", "OWASP", "OWASP Top 10 for Large Language Model Applications", "https://owasp.org/www-project-top-10-for-large-language-model-applications/",
         "AI Security", "Security Guide", "Intermediate to Advanced", "LLM basics", "Phase 11", 6, "Free", "Essential",
         "Prompt injection, insecure output handling, training data poisoning, model denial of service, excessive agency.",
         "Audit your sales agent against all 10 vulnerabilities and implement sanitization & tool permission boundaries.", "Not Started", "Non-negotiable security requirements before exposing agents to commercial clients.")
    ]

    for idx, r in enumerate(resources, start=5):
        ws.row_dimensions[idx].height = 26
        fill = fill_zebra if idx % 2 == 0 else fill_white
        for c_idx, val in enumerate(r, start=1):
            cell = ws.cell(row=idx, column=c_idx, value=val)
            cell.font = font_bold if c_idx in [1, 2, 3, 11, 12, 15] else font_data
            cell.fill = fill
            cell.border = border_thin
            
            if c_idx in [1, 2, 5, 6, 7, 9, 11, 12, 15]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif c_idx == 10:
                cell.alignment = Alignment(horizontal='right', vertical='center')
                cell.number_format = '#,##0'
            elif c_idx == 4:
                cell.alignment = Alignment(horizontal='left', vertical='center')
                if val and str(val).startswith("http"):
                    cell.font = font_link
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

    # Validations
    tot_row = len(resources) + 5
    dv_status = DataValidation(type="list", formula1='"Not Started,In Progress,Completed,Bookmarked"', allow_blank=True)
    ws.add_data_validation(dv_status)
    dv_status.add(f"O5:O{tot_row-1}")

    dv_imp = DataValidation(type="list", formula1='"Essential,Recommended,Optional"', allow_blank=True)
    ws.add_data_validation(dv_imp)
    dv_imp.add(f"L5:L{tot_row-1}")

    dv_free = DataValidation(type="list", formula1='"Free,Free / Freemium,Paid API,Free (Compute Req)"', allow_blank=True)
    ws.add_data_validation(dv_free)
    dv_free.add(f"K5:K{tot_row-1}")

    # Widths
    widths = [12, 16, 28, 25, 16, 16, 14, 18, 14, 12, 14, 13, 28, 28, 14, 25]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.auto_filter.ref = f"A4:P{tot_row-1}"

print("Official Resources module ready.")
