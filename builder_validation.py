import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

def build_sheet_validation(wb, font_family="Segoe UI"):
    ws = wb.create_sheet("Commercial Product Validation")
    ws.views.sheetView[0].showGridLines = True
    
    font_title = Font(name=font_family, size=14, bold=True, color="1A365D")
    font_sub = Font(name=font_family, size=10, italic=True, color="4A5568")
    font_header = Font(name=font_family, size=9, bold=True, color="FFFFFF")
    font_data = Font(name=font_family, size=9, color="2D3748")
    font_bold = Font(name=font_family, size=9, bold=True, color="2D3748")
    
    fill_header = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    border_thin = Border(
        left=Side(style='thin', color='CBD5E0'), right=Side(style='thin', color='CBD5E0'),
        top=Side(style='thin', color='CBD5E0'), bottom=Side(style='thin', color='CBD5E0')
    )

    # Title
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 18
    ws['A1'] = "COMMERCIAL PRODUCT & MARKET VALIDATION TRACKER"
    ws['A1'].font = font_title
    ws['A2'] = "Customer Discovery, Willingness to Pay, Unit Economics & Pilot Verification for Autonomous Commercial AI Employees"
    ws['A2'].font = font_sub

    headers = [
        "Opportunity ID", "Product / Agent Concept", "Target Industry",
        "Customer Segment", "Customer Core Problem", "Existing Manual Workflow",
        "Current Cost / Hours", "Proposed Autonomous AI Solution", "Interview Date",
        "Customer Feedback Summary", "Willingness to Pay ($/mo)", "Competitors & Alternatives",
        "Technical Feasibility", "Integration Requirements", "Est. Operating Cost ($/mo)",
        "Expected Business ROI", "Validation Status", "Next Action"
    ]
    
    header_row = 4
    ws.row_dimensions[header_row].height = 28
    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = border_thin
    ws.freeze_panes = "A5"

    opportunities = [
        (
            "VAL-001", "Autonomous Inbound SDR for B2B SaaS",
            "B2B SaaS / Tech Startups", "Mid-market sales directors & founders (10-50 employees)",
            "Leads who sign up on website wait 4-24 hours for follow-up; 50% drop off to competitors.",
            "Junior SDRs manually review form submissions, research company on LinkedIn, and send template emails.",
            "$4,500/mo salary per junior SDR + $500/mo tooling; 35 hours/week spent on manual qualification.",
            "Autonomous AI Sales Employee that responds within 60 seconds, qualifies via interactive chat/email, and books demos directly on rep calendars.",
            "2027-02-15", "Founders enthusiastic about sub-minute response time; primary concern is AI hallucinating features or making commitments.",
            "$499 - $1,200 / month", "Drift, Qualified, Intercom Fin, human SDR agencies",
            "High", "HubSpot CRM, Google Calendar, Slack webhooks, SendGrid", "$25 - $45 / month (tokens & cloud)",
            "3x increase in qualified demo bookings; pays for itself within first closed customer.",
            "Solution Validated", "Build working interactive prototype for 3 pilot SaaS companies."
        ),
        (
            "VAL-002", "Real Estate WhatsApp Inbound Qualifier & Booking Agent",
            "Real Estate Agencies", "Independent brokerage owners & top-producing real estate agents",
            "Brokers receive 100+ WhatsApp inquiries daily from portal listings; 70% are unqualified browsers, wasting hours of agent time.",
            "Agents manually reply to each WhatsApp message, ask budget/location questions, and send PDF brochures.",
            "$2,500/mo admin assistant salary; agents lose high-value buyers while chatting with unqualified renters.",
            "WhatsApp AI Sales Concierge that engages immediately, verifies buyer pre-approval & budget, shares matching listings, and schedules property viewings.",
            "2027-04-10", "Brokers love WhatsApp integration; must strictly avoid answering legal/title questions and hand off to licensed agents.",
            "$299 - $600 / month + $10/booked tour", "ManyChat, standard WhatsApp bots, human VAs",
            "High", "Meta WhatsApp Cloud API, Property MLS database / Airtable, Google Calendar", "$30 - $60 / month",
            "Saves 25 agent hours weekly and ensures zero missed leads on weekends and evenings.",
            "Problem Confirmed", "Connect LangGraph agent to WhatsApp sandbox with sample property catalog."
        ),
        (
            "VAL-003", "Solar & Home Services Appointment Setting Agent",
            "Residential Solar & HVAC", "Operations managers & sales leaders at residential solar contractors",
            "Homeowners request quotes online but ignore phone calls; cost per lead is $150+, and uncontacted leads burn marketing spend.",
            "Call center reps dial leads 5-8 times across 2 weeks; average contact rate is under 18%.",
            "$30,000/mo call center overhead; $150 customer acquisition cost per raw uncontacted lead.",
            "Omnichannel AI agent initiating instant SMS/WhatsApp and email follow-up sequence, calculating roof solar viability via RAG, and booking on-site consultations.",
            "2027-06-05", "Extreme urgency: lead acquisition costs are rising; converting 5% more leads doubles net profitability.",
            "$800 - $2,000 / month (or $50 / completed appointment)", "Podium, Verse.io, offshore call centers",
            "High", "Twilio SMS, Google Maps API, Jobber / ServiceTitan CRM", "$50 - $90 / month",
            "Increases lead contact rate from 18% to 45%; 200%+ ROI in month 1.",
            "Customer Interviewed", "Draft technical scope integrating Google Solar API with WhatsApp agent."
        ),
        (
            "VAL-004", "E-Commerce Cart Abandonment & Negotiation Concierge",
            "Direct-to-Consumer (DTC) E-Commerce", "E-commerce brand owners with average order value (AOV) > $150",
            "Shoppers abandon carts on high-ticket items ($200-$1,000) due to shipping costs, product questions, or indecision.",
            "Generic automated 10% discount email sequence sent 24 hours later; low engagement rate (8-12%).",
            "Loss of $50,000+ in abandoned cart revenue monthly for average store.",
            "AI Sales Employee reaching out via SMS/WhatsApp within 15 minutes, answering product sizing/compatibility questions, and offering dynamic authorized discounts.",
            "2027-07-12", "E-commerce owners want tight guardrails on discounts and easy Shopify App Store installation.",
            "$399 / month + 3% revenue share on recovered sales", "Klaviyo SMS, Gorgias, Cartloop",
            "Medium", "Shopify API, Klaviyo, WhatsApp / Twilio SMS", "$40 - $70 / month",
            "Recovers an extra $8,000 - $15,000 in monthly sales with zero staff intervention.",
            "Idea", "Build Shopify mock webhook listening for `carts/create` event."
        ),
        (
            "VAL-005", "IT Services Scope & Proposal Drafting Agent",
            "B2B IT & Software Agencies", "Agency founders & technical account executives (5-30 developers)",
            "Scoping customer custom software requests and drafting formal 15-page PDF proposals takes 10-15 hours per lead.",
            "Senior engineers and founders spend weekend hours writing scope documents, deliverables lists, and rate cards.",
            "$150/hr senior engineering opportunity cost; $1,500+ cost per proposal drafted.",
            "AI Sales Engineer that analyzes client requirements transcripts, drafts complete technical specification, line-item pricing, and generates branded PDF quote.",
            "2027-08-20", "Agencies excited about saving founder time; must allow easy human review and line-item adjustment before client receives PDF.",
            "$499 - $999 / month", "Proposify, PandaDoc, manual Google Docs",
            "High", "FastAPI backend, ReportLab PDF, OpenAI Structured Outputs, Slack", "$35 - $50 / month",
            "Cuts proposal preparation time from 12 hours to 30 minutes; doubles quote turnaround speed.",
            "Solution Validated", "Implement PDF quotation tool with ReportLab and Jinja2 templates (PRJ-008)."
        ),
        (
            "VAL-006", "Automated Commercial Fleet & Freight Brokerage Assistant",
            "Logistics & Freight Brokerage", "Freight brokers & mid-size trucking fleet dispatchers",
            "Brokers spend hours answering carrier phone calls asking 'is this load still available?' and negotiating per-mile rates.",
            "Brokers manually answer phone calls and emails, check TMS database, and haggle rates within $0.20/mile buffer.",
            "$3,500/mo freight coordinator salary per seat; high turnover and stress.",
            "AI Dispatch Employee answering email & web freight inquiries, confirming truck capacity, negotiating rate within predefined buffer, and generating rate confirmations.",
            "2027-09-08", "High volume, but industry has legacy software; TMS integration via API or email parser is mandatory.",
            "$600 - $1,500 / month per dispatcher seat", "Truckstop.com, DAT loadboards, manual broker teams",
            "Medium", "Transportation Management System (TMS) API / Email webhooks, PDF parser", "$45 - $80 / month",
            "Allows a 5-broker office to handle double the freight volume without adding staff.",
            "Problem Confirmed", "Identify open TMS API sandbox for logistics load query integration."
        )
    ]

    for idx, o in enumerate(opportunities, start=5):
        ws.row_dimensions[idx].height = 50
        fill = fill_zebra if idx % 2 == 0 else fill_white
        for c_idx, val in enumerate(o, start=1):
            cell = ws.cell(row=idx, column=c_idx, value=val)
            cell.font = font_bold if c_idx in [1, 2, 11, 13, 17] else font_data
            cell.fill = fill
            cell.border = border_thin
            
            if c_idx in [1, 9, 13, 17]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

    tot_row = len(opportunities) + 5
    # Validations
    dv_status = DataValidation(type="list", formula1='"Idea,Customer Interviewed,Problem Confirmed,Solution Validated,Pilot Running,Paid Pre-order,Abandoned"', allow_blank=True)
    ws.add_data_validation(dv_status)
    dv_status.add(f"Q5:Q{tot_row-1}")

    dv_feas = DataValidation(type="list", formula1='"High,Medium,Low"', allow_blank=True)
    ws.add_data_validation(dv_feas)
    dv_feas.add(f"M5:M{tot_row-1}")

    widths = [12, 26, 20, 25, 28, 28, 20, 32, 12, 28, 18, 22, 14, 25, 16, 24, 16, 26]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.auto_filter.ref = f"A4:R{tot_row-1}"

print("Commercial Product Validation module ready.")
