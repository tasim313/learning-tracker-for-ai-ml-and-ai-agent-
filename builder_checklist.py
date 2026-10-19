import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

def build_sheet_checklist(wb, font_family="Segoe UI"):
    ws = wb.create_sheet("Topic Completion Checklist")
    ws.views.sheetView[0].showGridLines = True
    
    font_title = Font(name=font_family, size=14, bold=True, color="1A365D")
    font_sub = Font(name=font_family, size=10, italic=True, color="4A5568")
    font_header = Font(name=font_family, size=9, bold=True, color="FFFFFF")
    font_data = Font(name=font_family, size=9, color="2D3748")
    font_bold = Font(name=font_family, size=9, bold=True, color="2D3748")
    
    fill_header = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    border_thin = Border(
        left=Side(style='thin', color='CBD5E0'), right=Side(style='thin', color='CBD5E0'),
        top=Side(style='thin', color='CBD5E0'), bottom=Side(style='thin', color='CBD5E0')
    )

    # Title
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 18
    ws['A1'] = "AI/ML & AI AGENT TOPIC COMPLETION CHECKLIST"
    ws['A1'].font = font_title
    ws['A2'] = "Granular Skill-by-Skill Verification: Python, Mathematics, ML, DL, LLMs, RAG, Agents, Production, and Business"
    ws['A2'].font = font_sub

    headers = [
        "Topic ID", "Subject", "Topic Name", "Priority", "Prerequisites",
        "Resource Links", "Status", "Confidence Score (1-5)", "Practice Task",
        "Assessment Criteria", "Evidence", "Date Completed", "Notes"
    ]
    
    header_row = 4
    ws.row_dimensions[header_row].height = 28
    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = border_thin
    ws.freeze_panes = "A5"

    topics = [
        # PYTHON
        ("TOP-001", "Python", "Variables & Advanced Data Types", "High", "Basic syntax", "RES-017", "Not Started", 3, "Implement custom cache dict with TTL expiration using Python dict & time.", "Explain memory references vs copies and mutability nuances.", "", "", ""),
        ("TOP-002", "Python", "Functions, Closures & Decorators", "High", "TOP-001", "RES-017", "Not Started", 3, "Build `@retry_with_backoff(max_retries=3)` decorator for API calls.", "Decorator properly wraps async/sync functions with functools.wraps.", "", "", ""),
        ("TOP-003", "Python", "OOP, Protocols & Abstract Classes", "High", "TOP-002", "RES-017", "Not Started", 3, "Design abstract BaseAgent class and implement EmailSalesAgent subclass.", "Demonstrate polymorphism, dependency injection, and clean interface segregation.", "", "", ""),
        ("TOP-004", "Python", "Exception Handling & Custom Errors", "High", "TOP-003", "RES-017", "Not Started", 3, "Create hierarchy: AgentError -> ToolExecutionError -> RateLimitExceededError.", "Code handles specific exceptions cleanly without catching broad Exception.", "", "", ""),
        ("TOP-005", "Python", "JSON, Serialization & Pydantic", "High", "TOP-001", "RES-017", "Not Started", 3, "Build Pydantic schema for LeadProfile with email and phone validator.", "Validates nested schemas and rejects malformed types gracefully.", "", "", ""),
        ("TOP-006", "Python", "AsyncIO & Concurrent HTTP Fetching", "High", "TOP-002", "RES-022", "Not Started", 2, "Use httpx.AsyncClient to query 20 website homepages concurrently.", "Runs without blocking event loop; handles timeouts and connection errors.", "", "", ""),
        ("TOP-007", "Python", "File I/O & JSONL Streaming", "Medium", "TOP-001", "RES-017", "Not Started", 3, "Stream 100,000 JSON lines without loading entire file into memory.", "Uses generator functions yielding records one line at a time.", "", "", ""),
        ("TOP-008", "Python", "Virtual Environments & Poetry/Pip", "High", "None", "RES-017", "Not Started", 3, "Configure clean Python 3.11 virtual environment with lockfile.", "Installs cleanly on clean Ubuntu machine without dependency conflicts.", "", "", ""),
        ("TOP-009", "Python", "Unit Testing with Pytest", "High", "TOP-003", "RES-017", "Not Started", 3, "Write 10 unit tests mocking external LLM responses with pytest-mock.", "All tests pass with > 85% branch coverage.", "", "", ""),
        ("TOP-010", "Python", "Structured Logging & Git Discipline", "Medium", "None", "RES-017", "Not Started", 3, "Configure Python `logging` to output JSON formatted logs with timestamps.", "Git repo has clean conventional commits and .gitignore for env/keys.", "", "", ""),

        # MATHEMATICS
        ("TOP-011", "Mathematics", "Vectors, Norms & Dot Products", "High", "High school math", "RES-018", "Not Started", 2, "Write Python function calculating Euclidean L2 norm and Dot Product.", "Explain geometric meaning of dot product (projection and angle alignment).", "", "", ""),
        ("TOP-012", "Mathematics", "Matrix Operations & Inversion", "High", "TOP-011", "RES-018", "Not Started", 2, "Multiply (4x3) matrix by (3x2) matrix by hand and in Python.", "Understand row-column compatibility and why inversion can be unstable.", "", "", ""),
        ("TOP-013", "Mathematics", "Probability & Bayes' Theorem", "High", "High school math", "RES-001", "Not Started", 2, "Calculate P(Lead Purchases | Visited Pricing Page) using Bayes' formula.", "Explain prior, likelihood, and posterior probability concepts.", "", "", ""),
        ("TOP-014", "Mathematics", "Descriptive Statistics & Distributions", "Medium", "TOP-013", "RES-001", "Not Started", 2, "Plot normal vs log-normal distribution of sales deal sizes in Matplotlib.", "Explain mean vs median sensitivity to outliers and standard deviation.", "", "", ""),
        ("TOP-015", "Mathematics", "Derivatives, Gradients & Chain Rule", "High", "Calculus basics", "RES-001", "Not Started", 2, "Calculate partial derivatives of MSE loss function with respect to weights.", "Deconstruct composite function derivatives using chain rule step-by-step.", "", "", ""),
        ("TOP-016", "Mathematics", "Gradient Descent Optimization", "High", "TOP-015", "RES-001", "Not Started", 2, "Implement Batch and Stochastic Gradient Descent in Python from scratch.", "Explain learning rate overshoot, oscillations, and convergence criteria.", "", "", ""),

        # MACHINE LEARNING
        ("TOP-017", "Machine Learning", "Linear & Polynomial Regression", "High", "TOP-016", "RES-016", "Not Started", 2, "Predict sales revenue from marketing spend and deal size with Scikit-Learn.", "Interpret regression coefficients, R^2, and residual plots.", "", "", ""),
        ("TOP-018", "Machine Learning", "Logistic Regression & Sigmoid", "High", "TOP-017", "RES-016", "Not Started", 2, "Train binary classifier predicting lead purchase probability.", "Explain sigmoid curve mapping (-inf, +inf) to [0, 1] probability.", "", "", ""),
        ("TOP-019", "Machine Learning", "Decision Trees & Random Forests", "High", "TOP-018", "RES-016", "Not Started", 2, "Train Random Forest and extract top-5 predictive lead features.", "Explain Gini impurity, entropy, bagging, and out-of-bag error.", "", "", ""),
        ("TOP-020", "Machine Learning", "Gradient Boosting with XGBoost", "High", "TOP-019", "RES-016", "Not Started", 2, "Train XGBoost model on customer churn tabular dataset.", "Explain boosting mechanism (fitting trees on residual errors sequentially).", "", "", ""),
        ("TOP-021", "Machine Learning", "Clustering with K-Means & PCA", "Medium", "TOP-012", "RES-016", "Not Started", 2, "Cluster sales prospects into 4 behavioral segments and visualize with PCA.", "Select optimal k using the elbow method and silhouette score.", "", "", ""),
        ("TOP-022", "Machine Learning", "Feature Engineering & Preprocessing", "High", "TOP-017", "RES-016", "Not Started", 2, "Build Scikit-Learn Pipeline with OneHotEncoder, StandardScaler, and Imputer.", "Pipeline runs transform on test set without leaking train statistics.", "", "", ""),
        ("TOP-023", "Machine Learning", "Train/Test Split & K-Fold CV", "High", "TOP-022", "RES-016", "Not Started", 2, "Implement 5-fold Stratified Cross-Validation on imbalanced sales dataset.", "Verify test split distribution matches real-world production ratio.", "", "", ""),
        ("TOP-024", "Machine Learning", "Overfitting, Regularization & Leakage", "High", "TOP-023", "RES-016", "Not Started", 2, "Diagnose overfitting via train vs validation loss curves; apply L1/L2 penalties.", "Identify 3 subtle data leakage traps in sales pipeline data.", "", "", ""),
        ("TOP-025", "Machine Learning", "Evaluation Metrics: Precision, Recall, ROC", "High", "TOP-024", "RES-016", "Not Started", 2, "Plot Precision-Recall curve and choose threshold maximizing high-value lead ROI.", "Explain why accuracy is misleading when positive class is only 2% of data.", "", "", ""),

        # DEEP LEARNING
        ("TOP-026", "Deep Learning", "PyTorch Tensors, Shapes & CUDA", "High", "TOP-012", "RES-014", "Not Started", 2, "Perform tensor operations on GPU; reshape with view() vs reshape().", "Manage memory and device transfers (cpu vs cuda) without errors.", "", "", ""),
        ("TOP-027", "Deep Learning", "Multi-Layer Perceptron (MLP)", "High", "TOP-026", "RES-014", "Not Started", 2, "Build PyTorch `nn.Module` with 2 hidden layers, ReLU, and Softmax.", "Model trains smoothly and achieves > 85% accuracy on benchmark data.", "", "", ""),
        ("TOP-028", "Deep Learning", "Activation & Loss Functions", "High", "TOP-027", "RES-014", "Not Started", 2, "Compare ReLU vs GELU vs LeakyReLU; CrossEntropyLoss vs BCEWithLogitsLoss.", "Explain dying ReLU problem and numerical stability of logit losses.", "", "", ""),
        ("TOP-029", "Deep Learning", "Backpropagation & Optimizers", "High", "TOP-028", "RES-014", "Not Started", 2, "Trace forward/backward gradient flow; compare SGD with momentum vs AdamW.", "Explain why AdamW decouples weight decay from gradient updates.", "", "", ""),
        ("TOP-030", "Deep Learning", "PyTorch Dataset & DataLoader", "High", "TOP-029", "RES-014", "Not Started", 2, "Build custom Dataset streaming sales conversation text and tabular metadata.", "Batching, shuffling, and worker multiprocessing execute without bottlenecks.", "", "", ""),

        # LLMS & TRANSFORMERS
        ("TOP-031", "LLMs", "Byte-Pair Encoding (BPE) Tokenization", "High", "TOP-001", "RES-012", "Not Started", 2, "Train mini BPE tokenizer on 500 sales emails using HuggingFace Tokenizers.", "Explain why numbers and non-English scripts take more tokens.", "", "", ""),
        ("TOP-032", "LLMs", "Self-Attention & Multi-Head Attention", "High", "TOP-026", "RES-013", "Not Started", 1, "Implement `MultiHeadAttention(nn.Module)` in pure PyTorch.", "QK^T / sqrt(d_k) equation correctly computed with causal attention mask.", "", "", ""),
        ("TOP-033", "LLMs", "Transformer Architecture Internals", "High", "TOP-032", "RES-013", "Not Started", 1, "Draw Transformer block diagram: LayerNorm, MHA, Residual, FeedForward.", "Explain encoder-only vs decoder-only architectures and their use cases.", "", "", ""),
        ("TOP-034", "LLMs", "Inference, KV Caching & Latency", "High", "TOP-033", "RES-009", "Not Started", 2, "Measure token generation speed with and without KV caching enabled.", "Explain memory consumption of KV cache across long multi-turn contexts.", "", "", ""),
        ("TOP-035", "LLMs", "Quantization: GGUF, AWQ, bitsandbytes", "High", "TOP-034", "RES-009", "Not Started", 2, "Run 4-bit quantized Llama-3 model locally using Ollama / llama.cpp.", "Explain precision loss trade-off between FP16, INT8, and INT4.", "", "", ""),

        # RAG (RETRIEVAL-AUGMENTED GENERATION)
        ("TOP-036", "RAG", "Chunking Strategies & Metadata", "High", "TOP-005", "RES-020", "Not Started", 2, "Chunk a 30-page software manual with recursive character and markdown splitting.", "Chunks preserve complete tables and include document title metadata.", "", "", ""),
        ("TOP-037", "RAG", "Embedding Models & Cosine Search", "High", "TOP-011", "RES-020", "Not Started", 2, "Embed 500 product FAQs with text-embedding-3-small and query top-3 matches.", "Explain embedding dimensions, normalization, and semantic drift.", "", "", ""),
        ("TOP-038", "RAG", "Vector Databases & PostgreSQL pgvector", "High", "TOP-037", "RES-023", "Not Started", 2, "Set up PostgreSQL table with `vector(1536)` and HNSW index in Docker.", "Query executes in < 15ms combining SQL WHERE clause with vector similarity.", "", "", ""),
        ("TOP-039", "RAG", "Hybrid Search & Cross-Encoder Rerank", "High", "TOP-038", "RES-020", "Not Started", 2, "Build pipeline combining BM25 keyword matching with dense vector reranking.", "Recall increases by > 15% on queries with specific product part numbers.", "", "", ""),
        ("TOP-040", "RAG", "Citations, Grounding & RAG Evaluation", "High", "TOP-039", "RES-004", "Not Started", 2, "Force model to output bracketed citations [Doc 1, p. 4] and test with RAGAS.", "Faithfulness score > 0.90; model refuses questions outside context.", "", "", ""),

        # AI AGENTS
        ("TOP-041", "AI Agents", "Tool Calling & Parameter Schemas", "High", "TOP-005", "RES-006", "Not Started", 2, "Write Pydantic tool schemas for CRM querying, email sending, calendar booking.", "LLM invokes tool with exact parameter types matching schema.", "", "", ""),
        ("TOP-042", "AI Agents", "ReAct Execution Loop", "High", "TOP-041", "RES-020", "Not Started", 2, "Code autonomous ReAct loop executing tools and feeding observations back.", "Loop handles tool execution failures and halts when final answer reached.", "", "", ""),
        ("TOP-043", "AI Agents", "Conversation State & Persistence", "High", "TOP-042", "RES-020", "Not Started", 2, "Persist multi-turn conversation graph state in SQLite / PostgreSQL database.", "Agent seamlessly resumes conversation context after application server reboot.", "", "", ""),
        ("TOP-044", "AI Agents", "State Graphs with LangGraph", "High", "TOP-043", "RES-020", "Not Started", 2, "Build LangGraph state graph with conditional edges routing based on lead intent.", "State transitions execute deterministically without deadlocks.", "", "", ""),
        ("TOP-045", "AI Agents", "Human-in-the-Loop (HITL) Checkpoints", "High", "TOP-044", "RES-020", "Not Started", 2, "Add interrupt gate requiring human manager approval for discounts > 10%.", "Workflow halts, serializes state, and resumes upon human API callback.", "", "", ""),

        # ADVANCED AI & PRODUCTION
        ("TOP-046", "Advanced AI", "LoRA & QLoRA Fine-Tuning", "High", "TOP-035", "RES-011", "Not Started", 1, "Fine-tune Llama-3-8B adapter using QLoRA with Unsloth on Colab/GPU.", "Training loss drops steadily; model adopts custom consultative sales tone.", "", "", ""),
        ("TOP-047", "Production", "FastAPI Asynchronous Microservice", "High", "TOP-006", "RES-022", "Not Started", 2, "Build FastAPI app with async `/chat` endpoint calling agent state graph.", "Handles 50 concurrent requests with response streaming and OpenAPI docs.", "", "", ""),
        ("TOP-048", "Production", "Docker Containerization & Compose", "High", "TOP-047", "RES-024", "Not Started", 2, "Write production Dockerfile and `docker-compose.yml` for FastAPI + Postgres.", "Containers build with minimal image size (< 500MB) running as non-root user.", "", "", ""),
        ("TOP-049", "Production", "LLM Security, Injection & Guardrails", "High", "TOP-045", "RES-025", "Not Started", 2, "Implement NeMo Guardrails / Llama-Guard blocking prompt injection attempts.", "Agent blocks 100% of tested jailbreak prompts attempting to alter prices.", "", "", ""),

        # BUSINESS & COMMERCIALIZATION
        ("TOP-050", "Business", "B2B Lead Qualification & BANT Logic", "High", "TOP-005", "RES-020", "Not Started", 2, "Define formal BANT scoring criteria and integrate into consultative prompt.", "Agent accurately categorizes leads into Hot, Warm, Cold with clear rationale.", "", "", ""),
        ("TOP-051", "Business", "Autonomous Price Negotiation Boundaries", "High", "TOP-045", "RES-007", "Not Started", 2, "Implement strict discount curve: 0-10% automatic, 11-20% manager approved, >20% forbidden.", "Agent never violates maximum authorized discount in 100 test simulations.", "", "", ""),
        ("TOP-052", "Business", "Commercial SaaS Pricing & Unit Economics", "High", "None", "RES-020", "Not Started", 2, "Calculate token API costs, server costs, and target 80% gross profit margins.", "Establish sustainable pricing: e.g. $499/mo per autonomous sales seat.", "", "", "")
    ]

    for idx, t in enumerate(topics, start=5):
        ws.row_dimensions[idx].height = 24
        fill = fill_zebra if idx % 2 == 0 else fill_white
        for c_idx, val in enumerate(t, start=1):
            cell = ws.cell(row=idx, column=c_idx, value=val)
            cell.font = font_bold if c_idx in [1, 3, 4, 7] else font_data
            cell.fill = fill
            cell.border = border_thin
            
            if c_idx in [1, 2, 4, 5, 6, 7, 8, 12]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

    tot_row = len(topics) + 5
    # Validations
    dv_status = DataValidation(type="list", formula1='"Not Started,In Progress,Completed,Mastered,Needs Revision,Blocked"', allow_blank=True)
    ws.add_data_validation(dv_status)
    dv_status.add(f"G5:G{tot_row-1}")

    dv_priority = DataValidation(type="list", formula1='"High,Medium,Low"', allow_blank=True)
    ws.add_data_validation(dv_priority)
    dv_priority.add(f"D5:D{tot_row-1}")

    dv_conf = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True)
    ws.add_data_validation(dv_conf)
    dv_conf.add(f"H5:H{tot_row-1}")

    widths = [11, 16, 26, 11, 14, 12, 14, 12, 32, 30, 20, 13, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.auto_filter.ref = f"A4:M{tot_row-1}"

print("Topic Completion Checklist module ready.")
