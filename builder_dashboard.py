import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference, Series

def build_sheet_dashboard(wb, font_family="Segoe UI"):
    ws = wb.create_sheet("Progress Dashboard")
    ws.views.sheetView[0].showGridLines = True
    
    font_title = Font(name=font_family, size=15, bold=True, color="1A365D")
    font_sub = Font(name=font_family, size=10, italic=True, color="4A5568")
    font_sec = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    font_sec_dark = Font(name=font_family, size=11, bold=True, color="1A365D")
    font_header = Font(name=font_family, size=9, bold=True, color="FFFFFF")
    font_data = Font(name=font_family, size=9, color="2D3748")
    font_bold = Font(name=font_family, size=9, bold=True, color="2D3748")
    font_kpi_num = Font(name=font_family, size=18, bold=True, color="1A365D")
    font_kpi_lbl = Font(name=font_family, size=9, bold=True, color="4A5568")
    
    fill_sec = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    fill_sec_accent = PatternFill(start_color="2C7A7B", end_color="2C7A7B", fill_type="solid")
    fill_header = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_kpi = PatternFill(start_color="F0F4F8", end_color="F0F4F8", fill_type="solid")
    fill_total = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")
    
    border_thin = Border(
        left=Side(style='thin', color='CBD5E0'), right=Side(style='thin', color='CBD5E0'),
        top=Side(style='thin', color='CBD5E0'), bottom=Side(style='thin', color='CBD5E0')
    )
    border_kpi = Border(
        left=Side(style='thin', color='CBD5E0'), right=Side(style='thin', color='CBD5E0'),
        top=Side(style='medium', color='1A365D'), bottom=Side(style='thin', color='CBD5E0')
    )
    border_double_bottom = Border(
        left=Side(style='thin', color='CBD5E0'), right=Side(style='thin', color='CBD5E0'),
        top=Side(style='thin', color='CBD5E0'), bottom=Side(style='double', color='1A365D')
    )

    # Title
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 18
    ws['A1'] = "EXECUTIVE AI LEARNING & PROJECT PERFORMANCE DASHBOARD"
    ws['A1'].font = font_title
    ws['A2'] = "Automated Real-Time Analytics Powered by Dynamic Formulas Connecting All 12 Curricula and Projects"
    ws['A2'].font = font_sub

    # Section 1 Header: Key Performance Indicators
    ws.row_dimensions[4].height = 22
    ws.merge_cells("A4:K4")
    c_sec1 = ws.cell(row=4, column=1, value="  1. STRATEGIC METRICS & LEARNING PROGRESS KPIS")
    c_sec1.font = font_sec
    c_sec1.fill = fill_sec
    c_sec1.alignment = Alignment(horizontal='left', vertical='center')

    # KPI Helper
    def make_kpi(start_r, start_c, span_c, title, formula, num_fmt=None):
        end_c = start_c + span_c - 1
        ws.merge_cells(start_row=start_r, start_column=start_c, end_row=start_r, end_column=end_c)
        ws.merge_cells(start_row=start_r+1, start_column=start_c, end_row=start_r+1, end_column=end_c)
        
        lbl = ws.cell(row=start_r, column=start_c, value=title)
        lbl.font = font_kpi_lbl
        lbl.alignment = Alignment(horizontal='center', vertical='center')
        lbl.fill = fill_kpi
        
        val = ws.cell(row=start_r+1, column=start_c, value=formula)
        val.font = font_kpi_num
        val.alignment = Alignment(horizontal='center', vertical='center')
        val.fill = fill_kpi
        if num_fmt:
            val.number_format = num_fmt
            
        for r in range(start_r, start_r+2):
            for c in range(start_c, end_c+1):
                ws.cell(row=r, column=c).border = border_kpi if r == start_r else border_thin
                ws.cell(row=r, column=c).fill = fill_kpi

    # Row 1 of KPIs (rows 5-6)
    ws.row_dimensions[5].height = 18
    ws.row_dimensions[6].height = 28
    make_kpi(5, 1, 2, "TOTAL TOPICS", "=COUNTA('Topic Completion Checklist'!C5:C56)", "#,##0")
    make_kpi(5, 3, 2, "COMPLETED TOPICS", "=COUNTIF('Topic Completion Checklist'!G5:G56, \"Completed\") + COUNTIF('Topic Completion Checklist'!G5:G56, \"Mastered\")", "#,##0")
    make_kpi(5, 5, 2, "TOPICS IN PROGRESS", "=COUNTIF('Topic Completion Checklist'!G5:G56, \"In Progress\")", "#,##0")
    make_kpi(5, 7, 2, "BLOCKED TOPICS", "=COUNTIF('Topic Completion Checklist'!G5:G56, \"Blocked\")", "#,##0")
    make_kpi(5, 9, 3, "OVERALL COMPLETION %", "=IF(A6>0, C6/A6, 0)", "0.0%")

    # Row 2 of KPIs (rows 8-9)
    ws.row_dimensions[7].height = 8 # spacer
    ws.row_dimensions[8].height = 18
    ws.row_dimensions[9].height = 28
    make_kpi(8, 1, 2, "PLANNED STUDY HOURS", "=SUM('Master Roadmap'!H5:H16)", "#,##0")
    make_kpi(8, 3, 2, "ACTUAL STUDY HOURS", "=SUM('Master Roadmap'!I5:I16)", "#,##0.0")
    make_kpi(8, 5, 2, "PROJECTS COMPLETED", "=COUNTIF('AI Project Backlog'!K5:K14, \"Completed\")", "#,##0")
    make_kpi(8, 7, 2, "TUTORIALS COMPLETED", "=COUNTIF('Tutorial and Course Tracker'!L5:L20, \"Completed\")", "#,##0")
    make_kpi(8, 9, 3, "AVG CONFIDENCE (1-5)", "=AVERAGE('Topic Completion Checklist'!H5:H56)", "0.0")

    # Section 2 Header: Curriculum Breakdown by Phase
    ws.row_dimensions[11].height = 8 # spacer
    ws.row_dimensions[12].height = 22
    ws.merge_cells("A12:F12")
    c_sec2 = ws.cell(row=12, column=1, value="  2. 12-PHASE ROADMAP EXECUTION & HOUR TRACKING")
    c_sec2.font = font_sec
    c_sec2.fill = fill_sec_accent
    c_sec2.alignment = Alignment(horizontal='left', vertical='center')

    phase_headers = ["Phase #", "Phase Name", "Planned Hours", "Actual Hours", "Status", "Progress %"]
    ws.row_dimensions[13].height = 24
    for idx, ph in enumerate(phase_headers, start=1):
        cell = ws.cell(row=13, column=idx, value=ph)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border_thin

    # 12 Phases linked directly to Master Roadmap
    for p_idx in range(1, 13):
        r_num = 13 + p_idx
        ref_row = 4 + p_idx # row in Master Roadmap
        ws.row_dimensions[r_num].height = 20
        fill = fill_zebra if p_idx % 2 == 0 else fill_white
        
        # Phase num
        c1 = ws.cell(row=r_num, column=1, value=f"Phase {p_idx}")
        c1.alignment = Alignment(horizontal='center', vertical='center')
        
        # Phase name
        c2 = ws.cell(row=r_num, column=2, value=f"='Master Roadmap'!B{ref_row}")
        c2.alignment = Alignment(horizontal='left', vertical='center')
        
        # Planned hours
        c3 = ws.cell(row=r_num, column=3, value=f"='Master Roadmap'!H{ref_row}")
        c3.alignment = Alignment(horizontal='right', vertical='center')
        c3.number_format = '#,##0'
        
        # Actual hours
        c4 = ws.cell(row=r_num, column=4, value=f"='Master Roadmap'!I{ref_row}")
        c4.alignment = Alignment(horizontal='right', vertical='center')
        c4.number_format = '#,##0.0'
        
        # Status
        c5 = ws.cell(row=r_num, column=5, value=f"='Master Roadmap'!N{ref_row}")
        c5.alignment = Alignment(horizontal='center', vertical='center')
        
        # Progress %
        c6 = ws.cell(row=r_num, column=6, value=f"=IF(C{r_num}>0, D{r_num}/C{r_num}, 0)")
        c6.alignment = Alignment(horizontal='right', vertical='center')
        c6.number_format = '0.0%'

        for c_i in range(1, 7):
            cell = ws.cell(row=r_num, column=c_i)
            cell.font = font_bold if c_i in [1, 5] else font_data
            cell.fill = fill
            cell.border = border_thin

    # Phase Table Totals
    tot_r = 26
    ws.row_dimensions[tot_r].height = 22
    ws.merge_cells(f"A{tot_r}:B{tot_r}")
    c_tot_lbl = ws.cell(row=tot_r, column=1, value="TOTAL ROADMAP HOURS")
    c_tot_lbl.font = Font(name=font_family, size=9, bold=True, color="1A365D")
    c_tot_lbl.alignment = Alignment(horizontal='right', vertical='center')
    c_tot_lbl.fill = fill_total
    for c_i in range(1, 3):
        ws.cell(row=tot_r, column=c_i).border = border_double_bottom
        ws.cell(row=tot_r, column=c_i).fill = fill_total

    c_p_sum = ws.cell(row=tot_r, column=3, value=f"=SUM(C14:C25)")
    c_p_sum.font = Font(name=font_family, size=9, bold=True, color="1A365D")
    c_p_sum.fill = fill_total
    c_p_sum.border = border_double_bottom
    c_p_sum.alignment = Alignment(horizontal='right', vertical='center')
    c_p_sum.number_format = '#,##0'

    c_a_sum = ws.cell(row=tot_r, column=4, value=f"=SUM(D14:D25)")
    c_a_sum.font = Font(name=font_family, size=9, bold=True, color="1A365D")
    c_a_sum.fill = fill_total
    c_a_sum.border = border_double_bottom
    c_a_sum.alignment = Alignment(horizontal='right', vertical='center')
    c_a_sum.number_format = '#,##0.0'

    c_stat = ws.cell(row=tot_r, column=5, value="-")
    c_stat.alignment = Alignment(horizontal='center', vertical='center')
    c_stat.fill = fill_total
    c_stat.border = border_double_bottom

    c_pct_tot = ws.cell(row=tot_r, column=6, value=f"=IF(C{tot_r}>0, D{tot_r}/C{tot_r}, 0)")
    c_pct_tot.font = Font(name=font_family, size=9, bold=True, color="1A365D")
    c_pct_tot.fill = fill_total
    c_pct_tot.border = border_double_bottom
    c_pct_tot.alignment = Alignment(horizontal='right', vertical='center')
    c_pct_tot.number_format = '0.0%'

    # Add Openpyxl BarChart comparing Planned vs Actual Hours by Phase
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Planned vs. Actual Study Hours by Curriculum Phase"
    chart.y_axis.title = "Hours"
    chart.x_axis.title = "Phase"
    chart.width = 17
    chart.height = 13

    data = Reference(ws, min_col=3, min_row=13, max_col=4, max_row=25)
    cats = Reference(ws, min_col=1, min_row=14, max_row=25)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, "H12")

    # Column Widths
    ws.column_dimensions['A'].width = 11
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 14
    ws.column_dimensions['F'].width = 14
    ws.column_dimensions['G'].width = 4
    ws.column_dimensions['H'].width = 14
    ws.column_dimensions['I'].width = 14
    ws.column_dimensions['J'].width = 14
    ws.column_dimensions['K'].width = 14

print("Progress Dashboard module ready.")
