import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_sheet_experiments(wb, font_family="Segoe UI"):
    ws = wb.create_sheet("AI Experiment Log")
    ws.views.sheetView[0].showGridLines = True
    
    font_title = Font(name=font_family, size=14, bold=True, color="1A365D")
    font_sub = Font(name=font_family, size=10, italic=True, color="4A5568")
    font_header = Font(name=font_family, size=9, bold=True, color="FFFFFF")
    font_data = Font(name=font_family, size=9, color="2D3748")
    font_bold = Font(name=font_family, size=9, bold=True, color="2D3748")
    font_link = Font(name=font_family, size=9, color="0969DA", underline="single")
    
    fill_header = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    border_thin = Border(
        left=Side(style='thin', color='CBD5E0'), right=Side(style='thin', color='CBD5E0'),
        top=Side(style='thin', color='CBD5E0'), bottom=Side(style='thin', color='CBD5E0')
    )

    # Title
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 18
    ws['A1'] = "SCIENTIFIC AI EXPERIMENT & BENCHMARK LOG"
    ws['A1'].font = font_title
    ws['A2'] = "Rigorous Measurement of Prompts, Hyperparameters, Latency, Token Costs, Tool Calling Accuracy & Failure Modes"
    ws['A2'].font = font_sub

    headers = [
        "Experiment ID", "Date", "Research Question / Hypothesis", "Model / Provider",
        "Model Version", "Dataset / Test Suite", "Prompt / Config Version",
        "Parameters", "Evaluation Benchmark", "Primary Metric",
        "Latency (ms)", "Cost / Run ($)", "Compute / VRAM", "Failure Cases Observed",
        "Comparison Baseline", "Conclusion & Key Takeaway", "Next Experiment & Link"
    ]
    
    header_row = 4
    ws.row_dimensions[header_row].height = 28
    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = border_thin
    ws.freeze_panes = "A5"

    experiments = [
        (
            "EXP-001", "2026-12-18",
            "Does XGBoost with tuned hyperparams outperform Random Forest on B2B lead conversion?",
            "XGBoost vs Scikit-Learn RF", "xgboost v2.0.3, sklearn v1.4", "10,000 anonymized SaaS CRM records", "v1.0 (5-fold stratified CV)",
            "n_estimators=300, max_depth=5, lr=0.05, subsample=0.8", "Held-out test set (2,000 records)", "ROC-AUC: 0.864 vs 0.821",
            "145 ms (inference 100 rows)", "$0.00", "CPU 4 cores, 1.2 GB RAM",
            "Slight drop in recall on newly founded startups (< 6 months operating history).",
            "Scikit-Learn baseline RandomForest (AUC: 0.821)",
            "XGBoost improved ROC-AUC by 4.3 percentage points and inference was 2.5x faster.",
            "EXP-002: Add interaction terms between website visits and company headcount."
        ),
        (
            "EXP-002", "2027-01-22",
            "Can PyTorch MLP match XGBoost accuracy on tabular CRM data with learned categorical embeddings?",
            "PyTorch Tabular MLP", "PyTorch 2.2 CUDA 12.1", "10,000 SaaS CRM records with 4 categorical features", "TabularNet v1.2",
            "Embedding dim=16, 2 hidden layers (128, 64), AdamW lr=1e-3, batch=64", "Held-out test set (2,000 records)", "ROC-AUC: 0.835 vs 0.864",
            "85 ms (GPU batch)", "$0.00", "GPU RTX 3080, 800 MB VRAM",
            "Overfits faster than tree ensembles on sparse industry categories with low sample count.",
            "XGBoost baseline from EXP-001 (AUC: 0.864)",
            "XGBoost remains superior on purely tabular CRM data; deep learning requires 10x more data.",
            "EXP-003: Focus deep learning on text message intent classification where it excels."
        ),
        (
            "EXP-003", "2027-03-05",
            "Does OpenAI native Structured Outputs eliminate malformed JSON errors compared to zero-shot prompting?",
            "OpenAI GPT-4o-mini", "gpt-4o-mini-2024-07-18", "100 complex, messy sales call transcripts", "Prompt v2.1 with Pydantic BANT schema",
            "temperature=0.0, response_format=PydanticModel, max_tokens=1000", "100 transcript test suite", "Schema Validity: 100% vs 87%",
            "820 ms avg", "$0.00042 / call", "Cloud API",
            "Zero schema failures with native structured outputs. Zero-shot baseline failed on nested arrays.",
            "Zero-shot prompting with markdown ```json codeblocks",
            "Native Structured Outputs guarantees 100% schema compliance and completely eliminates retry loops.",
            "EXP-004: Test extraction latency and cost between GPT-4o-mini and Gemini-1.5-Flash."
        ),
        (
            "EXP-004", "2027-03-12",
            "Cost & Latency comparison: GPT-4o-mini vs Gemini-1.5-Flash for real-time lead qualification.",
            "GPT-4o-mini vs Gemini-1.5-Flash", "gpt-4o-mini vs gemini-1.5-flash-001", "50 customer chat dialogue transcripts", "Standardized qualification prompt v3.0",
            "temp=0.0, strict structured JSON output mode", "50 multi-turn test dialogues", "F1 on BANT extraction: 0.92 vs 0.90",
            "GPT-4o-mini: 680ms | Gemini: 420ms", "GPT: $0.00038 | Gemini: $0.00018", "Cloud API",
            "Gemini-1.5-Flash was occasionally too conservative on extracting numeric budget values.",
            "GPT-4o-mini baseline",
            "Gemini-1.5-Flash is 38% faster and 52% cheaper, making it the ideal choice for high-volume chat webhooks.",
            "EXP-005: Benchmark RAG chunk sizes for product catalog retrieval."
        ),
        (
            "EXP-005", "2027-04-10",
            "What is the optimal chunk size (256 vs 512 vs 1024 tokens) for product specification retrieval?",
            "OpenAI text-embedding-3-small + pgvector", "text-embedding-3-small, dim=1536", "150-page enterprise software catalog & pricing sheets", "RecursiveCharacterTextSplitter with overlap=10%",
            "Cosine similarity, top_k=5 retrieval", "50 synthetic product domain queries", "Recall@5: 256t: 74% | 512t: 91% | 1024t: 82%",
            "18 ms (PostgreSQL pgvector query)", "$0.00002 / query", "Local PostgreSQL container, 500 MB RAM",
            "256-token chunks fragmented pricing tables; 1024-token chunks diluted specific SKU keyword relevance.",
            "Naive 1000-character fixed chunking (Recall@5: 68%)",
            "512-token chunking with 50-token overlap yielded the highest retrieval accuracy (91% Recall@5).",
            "EXP-006: Add BM25 hybrid search and cross-encoder reranking to 512-token chunks."
        ),
        (
            "EXP-006", "2027-04-18",
            "Does Cross-Encoder reranking improve RAG retrieval accuracy on exact product part numbers?",
            "Cohere Rerank v3 vs Dense Cosine", "cohere/rerank-v3.5 + text-embedding-3-small", "Software catalog with 40 specific part numbers/SKUs", "Hybrid retrieval pipeline v2.0",
            "top_k_retrieve=20, top_n_rerank=5", "40 exact SKU and part number queries", "Top-1 Accuracy: 95% vs 72%",
            "120 ms (Dense: 22ms + Rerank: 98ms)", "$0.0012 / query", "API / Local container",
            "Minor latency increase (+98ms), well within interactive chat tolerance (< 1.5s).",
            "Dense vector search alone (72% Top-1 Accuracy)",
            "Cross-Encoder reranking drastically improves retrieval on specific technical terms and part numbers.",
            "EXP-007: Test tool calling reliability of Llama-3-8B locally vs GPT-4o-mini."
        ),
        (
            "EXP-007", "2027-05-18",
            "Can local Llama-3-8B-Instruct reliably call 4 business tools without argument hallucination?",
            "Meta Llama-3-8B-Instruct (4-bit)", "Llama-3-8B-Instruct-GGUF (Q4_K_M) via Ollama", "50 multi-step tool execution scenarios", "ReAct tool prompt with JSON schema definitions",
            "temperature=0.1, max_tokens=512, top_p=0.9", "50 simulated tool-use test tasks", "Tool Call Accuracy: 84% vs 98% (GPT-4o-mini)",
            "1,450 ms (local RTX 3080)", "$0.00 (Self-hosted)", "6.2 GB VRAM on GPU",
            "Llama-3-8B sometimes hallucinated optional parameters or invoked non-existent tools when confused.",
            "GPT-4o-mini tool execution accuracy (98%)",
            "Llama-3-8B is viable for single-tool tasks, but multi-tool agentic planning still requires frontier APIs.",
            "EXP-008: Test system prompt defenses against adversarial discount extraction attacks."
        ),
        (
            "EXP-008", "2027-08-10",
            "Can adversarial prompt injection bypass autonomous sales agent pricing guardrails?",
            "GPT-4o-mini with NeMo Guardrails", "gpt-4o-mini + NeMo Guardrails v0.8", "30 adversarial jailbreak prompts ('Forget previous instructions', etc.)", "System prompt v4.0 with deterministic code guardrails",
            "temperature=0.0, code-level price floor validator", "30 adversarial prompt attack suite", "Jailbreak Defense Rate: 100%",
            "710 ms avg", "$0.00045 / test", "Cloud API + local validator",
            "LLM was tricked twice by clever roleplay, BUT deterministic code layer caught and blocked the discount.",
            "Prompt-only defense (failed 6 out of 30 attacks; 80% defense rate)",
            "NEVER rely solely on prompt instructions for business security. Always enforce hard code-level bounds.",
            "Next: Deploy production platform with deterministic audit logger."
        )
    ]

    for idx, e in enumerate(experiments, start=5):
        ws.row_dimensions[idx].height = 45
        fill = fill_zebra if idx % 2 == 0 else fill_white
        for c_idx, val in enumerate(e, start=1):
            cell = ws.cell(row=idx, column=c_idx, value=val)
            cell.font = font_bold if c_idx in [1, 4, 10] else font_data
            cell.fill = fill
            cell.border = border_thin
            
            if c_idx in [1, 2, 4, 10]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif c_idx in [11, 12]:
                cell.alignment = Alignment(horizontal='right', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

    tot_row = len(experiments) + 5
    widths = [12, 12, 28, 22, 20, 24, 22, 25, 24, 22, 14, 14, 18, 28, 25, 28, 28]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.auto_filter.ref = f"A4:Q{tot_row-1}"

print("AI Experiment Log module ready.")
